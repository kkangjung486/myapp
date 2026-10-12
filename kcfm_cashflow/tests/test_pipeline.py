"""샘플 PDF → 파싱 → 표준화 → 계산 전 과정 테스트 (정답: sample_data.sample_company)."""
import pandas as pd
import pytest

from data_pipeline.aggregation import aggregate, to_values
from data_pipeline.edits import apply_edits
from data_pipeline.extraction import block_metas, build_rows, default_analysis_years, select_sources
from diagnostic_rules.rules import run_diagnostics
from financial_calculations.summary import analyze
from financial_calculations.validation import has_failure, run_checks
from pdf_parser.statement_parser import parse_pdf
from sample_data.generate_sample_pdfs import generate
from sample_data.sample_company import build

K = 1000  # 원천 데이터는 천원 단위


@pytest.fixture(scope="module")
def pdfs(tmp_path_factory):
    return generate(str(tmp_path_factory.mktemp("pdfs")))


@pytest.fixture(scope="module")
def pipeline(pdfs):
    docs = [parse_pdf(open(p, "rb").read(), p.split("/")[-1]) for p in pdfs if "3개년" not in p]
    metas = block_metas(docs)
    rows = build_rows(docs, metas, "제조업")
    selected, choices, warnings = select_sources(rows)
    return docs, metas, rows, selected, choices, warnings


def test_document_recognition(pipeline):
    docs, metas, *_ = pipeline
    assert all(d.company_name == "주식회사 케이씨정밀" for d in docs)
    units = {(m.file[:4], m.statement): m.unit for m in metas}
    assert units[("2023", "BS")] == "천원" and units[("2025", "BS")] == "원"
    assert {m.statement for m in metas if m.file.startswith("2025")} == {"BS", "IS", "COGS", "CF"}
    assert [m.years for m in metas if m.file.startswith("2024")][0] == [2024, 2023]


def test_prior_year_column_cross_check_after_unit_conversion(pipeline):
    """2024년이 '2024 PDF(천원)'와 '2025 PDF(원)'에 모두 있어도 환산 후 일치해야 한다."""
    *_, warnings = pipeline
    assert not [w for w in warnings if "대조 불일치" in w]


def test_current_year_source_preferred(pipeline):
    *_, choices, _ = pipeline
    c = {(c.statement, c.year): c for c in choices}
    assert c[("BS", 2024)].file.startswith("2024") and c[("BS", 2024)].col == 0
    assert c[("BS", 2022)].col == 1          # 2022는 2023 PDF의 전기 열
    assert c[("CF", 2024)].file.startswith("2025") and c[("CF", 2024)].col == 1


@pytest.mark.parametrize("year", [2022, 2023, 2024, 2025])
def test_standardized_values_match_ground_truth(pipeline, year):
    *_, selected, _, _ = pipeline
    v = aggregate(selected)[year].values
    t = build()[year]
    expected = {
        "CASH": t["cash"], "AR": t["ar_net"], "INV": t["inv"], "AP": t["ap"], "PPE": t["ppe"],
        "ST_BORROW": t["st_borrow"], "LT_BORROW": t["lt_borrow"], "CURRENT_LTD": t["cur_ltd"],
        "CA_TOTAL": t["ca"], "CL_TOTAL": t["cl"], "TOTAL_ASSETS": t["assets"], "TOTAL_LIAB": t["liab"],
        "TOTAL_EQUITY": t["equity"], "REVENUE": t["revenue"], "COGS": t["cogs"], "NET_INCOME": t["ni"],
        "DEPRECIATION": t["dep_total"], "PURCHASE_MERCH": t["merch_purchase"], "PURCHASE_RAW": t["raw_purchase"],
        "RETIREMENT": t["retire_net"],
    }
    for code, val in expected.items():
        assert v[code] == pytest.approx(val * K), code


def test_disclosed_cfo_and_validation(pipeline):
    _, metas, _, selected, *_ = pipeline
    std = aggregate(selected)
    truth = build()
    assert std[2025].values["CFO"] == pytest.approx(truth[2025]["cf"]["cfo"] * K)
    assert std[2024].values["CFO"] == pytest.approx(truth[2024]["cf"]["cfo"] * K)
    assert "CFO" not in std[2023].values
    years = default_analysis_years(selected)
    assert years == [2023, 2024, 2025]
    checks = run_checks(std, selected, years, metas)
    assert not has_failure(checks)
    bs = [c for c in checks if c.item == "자산총계 = 부채총계 + 자본총계"]
    assert len(bs) == 3 and all(c.status == "통과" for c in bs)


def test_analysis_and_diagnostics_on_sample(pipeline):
    *_, selected, _, _ = pipeline
    std = aggregate(selected)
    an = analyze(to_values(std), [2023, 2024, 2025])
    assert [a.cfo_basis for a in an] == ["추정", "공시", "공시"]
    assert an[0].ccc.ar_method.startswith("평균잔액(")       # 2022 전기 열 덕분에 평균잔액 사용
    titles = [s.title for s in run_diagnostics(an, "제조업")]
    assert any("흑자인데 영업현금흐름은 마이너스" in t for t in titles)
    assert any("CCC" in t and "연속 악화" in t for t in titles)


def test_three_years_in_one_pdf(pdfs):
    path = [p for p in pdfs if "3개년" in p][0]
    doc = parse_pdf(open(path, "rb").read(), "3개년.pdf")
    assert [b.years for b in doc.blocks] == [[2025, 2024, 2023], [2025, 2024, 2023]]
    metas = block_metas([doc])
    selected, *_ = select_sources(build_rows([doc], metas, "제조업"))
    std = aggregate(selected)
    truth = build()
    for y in (2023, 2024, 2025):
        assert std[y].values["TOTAL_ASSETS"] == pytest.approx(truth[y]["assets"] * K)
        assert std[y].values["NET_INCOME"] == pytest.approx(truth[y]["ni"] * K)
        assert std[y].values["INV"] == pytest.approx(truth[y]["inv"] * K)


def test_user_edits_are_tracked_separately(pipeline):
    *_, selected, _, _ = pipeline
    base = selected[selected["year"] == 2025].reset_index(drop=True)
    edited = base.copy()
    i = edited.index[edited["label"].str.contains("선급금")][0]
    edited.at[i, "std_label"] = "기타영업유동자산"
    j = edited.index[edited["label"].str.contains("현금및현금성자산")][0]
    edited.at[j, "confirmed_won"] = 1.0
    new = {c: None for c in edited.columns}
    new.update(label="수동 계정", confirmed_won=5_000_000.0, std_label="재고자산", include=True, statement_name="재무상태표")
    edited = pd.concat([edited, pd.DataFrame([new])], ignore_index=True)
    out = apply_edits(base, edited, 2025)
    assert out.at[i, "status"] == "수정됨"
    assert out.at[j, "status"] == "수정됨"
    assert out.at[j, "extracted_won"] == base.at[j, "extracted_won"]     # 원본 추출값 보존
    assert out.iloc[-1]["status"] == "수동 추가" and out.iloc[-1]["statement"] == "BS"
    v = aggregate(out)[2025].values
    assert v["OTHER_OP_CA"] == pytest.approx(55_000_000)
    assert v["INV"] == pytest.approx(build()[2025]["inv"] * K + 5_000_000)


def test_balance_sheet_failure_detected(pipeline):
    _, metas, _, selected, *_ = pipeline
    broken = selected.copy()
    idx = broken.index[(broken["year"] == 2025) & (broken["std_label"] == "자산총계")][0]
    broken.at[idx, "confirmed_won"] += 1_000_000
    checks = run_checks(aggregate(broken), broken, [2023, 2024, 2025], metas)
    assert has_failure(checks)


def test_duplicate_total_flagged():
    df = pd.DataFrame([
        {"year": 2025, "statement": "IS", "std_label": "매출액", "include": True, "confirmed_won": 100.0, "unit": "원"},
        {"year": 2025, "statement": "IS", "std_label": "매출액", "include": True, "confirmed_won": 120.0, "unit": "원"},
    ])
    sy = aggregate(df)[2025]
    assert sy.values["REVENUE"] == 100.0
    assert any("중복" in n for n in sy.notes["REVENUE"])


def test_excel_export_has_required_sheets(pipeline):
    import io

    from openpyxl import load_workbook

    from diagnostic_rules.rules import run_diagnostics
    from export.excel_export import build_excel
    from financial_calculations.validation import checks_frame

    docs, metas, _, selected, _, warnings = pipeline
    std = aggregate(selected)
    an = analyze(to_values(std), [2023, 2024, 2025])
    data = build_excel(company="주식회사 케이씨정밀", industry="제조업", analysis_years=[2023, 2024, 2025],
                       files=[d.filename for d in docs], metas=metas, warnings=warnings, selected_rows=selected,
                       std=std, analyses=an, signals=run_diagnostics(an, "제조업"),
                       checks_df=checks_frame(run_checks(std, selected, [2023, 2024, 2025], metas)))
    wb = load_workbook(io.BytesIO(data))
    for name in ["기업 기본정보", "3개년 추출 재무제표", "표준화 재무데이터", "현금흐름 분석", "운전자본 분석",
                 "CCC 분석", "진단 결과", "검증 결과", "계산식"]:
        assert name in wb.sheetnames
