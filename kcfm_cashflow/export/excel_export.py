"""분석 결과 Excel 다운로드."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.accounts import STANDARD_ACCOUNTS
from financial_calculations.summary import YearAnalysis, metrics_table

FORMULAS = [
    ("공시 영업활동현금흐름", "현금흐름표의 '영업활동으로 인한 현금흐름' 값 그대로 사용"),
    ("추정 영업활동현금흐름", "당기순이익 + 감가상각비 + 무형자산상각비 + 퇴직급여충당부채(순액) 증가 - 매출채권 증가 - 계약자산 증가 - 재고 증가 - 기타영업유동자산 증가 + 매입채무 증가 + 계약부채 증가 + 기타영업유동부채 증가 (재무제표만을 이용한 추정치)"),
    ("영업현금흐름률", "영업활동현금흐름 ÷ 매출액 × 100"),
    ("현금이익전환율", "영업활동현금흐름 ÷ 당기순이익 (순이익 0 이하이면 해석 불가)"),
    ("순운전자본(NWC)", "유동자산 - 유동부채"),
    ("영업운전자본(OWC)", "매출채권 + 계약자산 + 재고자산 + 선택된 기타영업유동자산 - 매입채무 - 계약부채 - 선택된 기타영업유동부채"),
    ("DSO", "평균 매출채권 ÷ 매출액 × 365"),
    ("DIO", "평균 재고자산 ÷ 매출원가 × 365"),
    ("DPO", "평균 매입채무 ÷ 연간 매입액 × 365 (매입액 없으면 매출원가 대용 – 표시됨)"),
    ("CCC", "DSO + DIO - DPO"),
    ("평균잔액", "(전기말 + 당기말) ÷ 2, 전기말 자료가 없으면 기말잔액 사용(표시됨)"),
    ("총차입금", "단기차입금 + 유동성장기부채 + 장기차입금 + 사채 (리스부채 제외)"),
    ("현금 증감 분해", "순이익 - 매출채권 증가 - 재고 증가 + 매입채무 증가 - 유형자산 장부가 증가 + 차입금 증가 + 기타(잔여) = 현금 증감"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_excel(*, company: str, industry: str, analysis_years: list[int], files: list[str],
                metas: list, warnings: list[str], selected_rows: pd.DataFrame, std: dict,
                analyses: list[YearAnalysis], signals: list, checks_df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        _info_sheet(xw, company, industry, analysis_years, files, metas, warnings, checks_df)
        _extracted_sheet(xw, selected_rows)
        _standard_sheet(xw, std)
        _cashflow_sheet(xw, analyses)
        _wc_sheet(xw, analyses)
        _ccc_sheet(xw, analyses)
        _diag_sheet(xw, signals)
        checks_df.to_excel(xw, sheet_name="검증 결과", index=False)
        pd.DataFrame(FORMULAS, columns=["지표", "계산식"]).to_excel(xw, sheet_name="계산식", index=False)
        for ws in xw.book.worksheets:
            _style(ws)
    return buf.getvalue()


def _info_sheet(xw, company, industry, years, files, metas, warnings, checks_df):
    fail = (checks_df["결과"] == "실패").any() if not checks_df.empty else False
    rows = [
        ("회사명", company), ("업종", industry),
        ("분석 연도", ", ".join(str(y) for y in years)),
        ("작성일시", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("업로드 파일", ", ".join(files)),
        ("검증 결과", "실패 항목 있음 – 숫자를 확정적으로 사용하지 마세요" if fail else "실패 항목 없음"),
        ("금액 단위", "모든 금액은 '원' 단위로 환산되어 있습니다."),
        ("도구", "KCFM 기업 현금흐름 진단 (MVP)"),
    ]
    df = pd.DataFrame(rows, columns=["항목", "내용"])
    df.to_excel(xw, sheet_name="기업 기본정보", index=False)
    meta_df = pd.DataFrame([{
        "파일": m.file, "재무제표": m.statement, "단위": m.unit, "단위 인식": m.unit_source,
        "연도(열 순서)": ", ".join(str(y) for y in m.years), "연도 인식": m.year_source,
        "형식": "연도별 2열" if m.layout == "paired" else "연도별 1열", "페이지": ", ".join(map(str, m.pages)),
    } for m in metas])
    meta_df.to_excel(xw, sheet_name="기업 기본정보", index=False, startrow=len(df) + 2)
    if warnings:
        pd.DataFrame({"인식 경고": warnings}).to_excel(xw, sheet_name="기업 기본정보", index=False,
                                                    startrow=len(df) + len(meta_df) + 5)


def _extracted_sheet(xw, rows: pd.DataFrame):
    if rows.empty:
        pd.DataFrame().to_excel(xw, sheet_name="3개년 추출 재무제표")
        return
    out = pd.DataFrame({
        "연도": rows["year"], "재무제표": rows["statement_name"], "원본 계정명": rows["label"],
        "원문 숫자": rows["raw_text"], "원문 단위": rows["unit"], "추출금액(원)": rows["extracted_won"],
        "확인금액(원)": rows["confirmed_won"],
        "수정 여부": (rows["confirmed_won"].fillna(-1e30) != rows["extracted_won"].fillna(-1e30)).map({True: "수정", False: ""}),
        "표준 계정": rows["std_label"], "분석 포함": rows["include"].map({True: "O", False: "X"}),
        "상태": rows["status"], "메모": rows["note"], "파일": rows["file"], "페이지": rows["page"],
    })
    out.to_excel(xw, sheet_name="3개년 추출 재무제표", index=False)


def _standard_sheet(xw, std: dict):
    years = sorted(std)
    recs = []
    for acc in STANDARD_ACCOUNTS:
        if acc.kind == "excluded":
            continue
        vals = {f"{y}년": std[y].values.get(acc.code) for y in years}
        if all(v is None for v in vals.values()):
            continue
        notes = "; ".join(n for y in years for n in std[y].notes.get(acc.code, []))
        recs.append({"표준 계정": acc.label, **vals, "비고": notes})
    pd.DataFrame(recs).to_excel(xw, sheet_name="표준화 재무데이터", index=False)


def _cashflow_sheet(xw, analyses):
    recs = []
    for a in analyses:
        recs.append({"연도": a.year, "항목": "공시 영업활동현금흐름", "금액(원)": a.cfo_disclosed, "설명": "현금흐름표 기재값" if a.cfo_disclosed is not None else "현금흐름표 없음"})
        est = a.cfo_estimate
        for c in est.components:
            recs.append({"연도": a.year, "항목": f"  [추정] {c.label}", "금액(원)": c.amount, "설명": c.rule})
        recs.append({"연도": a.year, "항목": "추정 영업활동현금흐름", "금액(원)": est.value, "설명": " / ".join(est.warnings)})
        recs.append({"연도": a.year, "항목": "공시 - 추정 차이", "금액(원)": a.cfo_diff, "설명": ""})
        recs.append({"연도": a.year, "항목": "분석에 사용한 영업활동현금흐름", "금액(원)": a.cfo, "설명": a.cfo_basis or "확인 필요"})
        recs.append({"연도": a.year, "항목": "영업현금흐름률(%)", "금액(원)": a.ocf_margin, "설명": ""})
        recs.append({"연도": a.year, "항목": "현금이익전환율(배)", "금액(원)": a.cash_conversion, "설명": a.cash_conversion_note})
        if a.bridge:
            for it in a.bridge.items:
                recs.append({"연도": a.year, "항목": f"[현금 증감 분해] {it.label}", "금액(원)": it.amount, "설명": ""})
            recs.append({"연도": a.year, "항목": "[현금 증감 분해] 최종 현금 증감", "금액(원)": a.bridge.cash_change,
                         "설명": " ".join(a.bridge.narrative)})
    pd.DataFrame(recs).to_excel(xw, sheet_name="현금흐름 분석", index=False)
    metrics_table(analyses).reset_index().rename(columns={"index": "지표"}).to_excel(
        xw, sheet_name="현금흐름 분석", index=False, startcol=6)


def _wc_sheet(xw, analyses):
    from config.accounts import label_of
    recs = []
    for a in analyses:
        recs.append({"연도": a.year, "항목": "순운전자본(NWC) = 유동자산 - 유동부채", "금액(원)": a.nwc})
        for code, v in a.owc_breakdown.items():
            if v:
                recs.append({"연도": a.year, "항목": f"  OWC 구성: {label_of(code)}", "금액(원)": v})
        recs.append({"연도": a.year, "항목": "영업운전자본(OWC)", "금액(원)": a.owc})
        recs.append({"연도": a.year, "항목": "전년 대비 OWC 증감", "금액(원)": a.owc_change})
        recs.append({"연도": a.year, "항목": "해석", "금액(원)": None, "해석": a.owc_change_text})
    pd.DataFrame(recs).to_excel(xw, sheet_name="운전자본 분석", index=False)


def _ccc_sheet(xw, analyses):
    recs = []
    for a in analyses:
        c = a.ccc
        recs.append({
            "연도": a.year, "DSO(일)": c.dso, "DIO(일)": c.dio, "DPO(일)": c.dpo, "CCC(일)": c.ccc,
            "평균 매출채권": c.avg_ar, "매출채권 잔액 기준": c.ar_method,
            "평균 재고자산": c.avg_inv, "재고 잔액 기준": c.inv_method,
            "평균 매입채무": c.avg_ap, "매입채무 잔액 기준": c.ap_method,
            "DPO 분모(매입액)": c.purchases, "DPO 분모 기준": c.purchases_basis,
            "참고 DPO(재고증감 반영 추정매입액)": c.dpo_ref, "참고 DSO(계약자산 포함)": c.dso_with_contract,
            "비고": " / ".join(c.notes),
        })
    pd.DataFrame(recs).to_excel(xw, sheet_name="CCC 분석", index=False)


def _diag_sheet(xw, signals):
    pd.DataFrame([{"신호": s.level, "제목": s.title, "내용": s.detail, "관련 연도": ", ".join(map(str, s.years)),
                   "영업CF 기준": s.basis} for s in signals]
                 or [{"신호": "-", "제목": "탐지된 신호 없음", "내용": "", "관련 연도": "", "영업CF 기준": ""}]
                 ).to_excel(xw, sheet_name="진단 결과", index=False)


def _style(ws):
    for cell in ws[1]:
        if cell.value is not None:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), start=1):
        width = 10
        for cell in col:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                if abs(cell.value) >= 1000:
                    cell.number_format = "#,##0;[Red](#,##0)"
                else:
                    cell.number_format = "#,##0.0;[Red](#,##0.0)"
                width = max(width, 16)
            elif cell.value is not None:
                width = max(width, min(60, len(str(cell.value)) * 1.6 + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
